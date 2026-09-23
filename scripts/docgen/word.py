"""
================================================================================
scripts/docgen/word.py — the Word (.docx) writer
================================================================================
A thin, opinionated layer over python-docx so every KairoX service document
comes out looking the same: same cover page, same heading colours, same table
style, same code-block shading.

It also renders a subset of Markdown, which is how the SYSTEM GUIDES are
authored: plain `.md` files under `docs/service/_src/` that a human can edit and
git can diff, turned into Word at build time. Supported:

    # ## ### ####      headings
    plain paragraphs   (blank-line separated)
    - item             bullet list (two-space indent = sub-bullet)
    1. item            numbered list
    | a | b |          table, with the --- separator row
    ```                fenced code block
    > text             call-out box
    ---                horizontal rule / spacer
    **bold**  `code`   inline

Nothing exotic — a frontend developer opening the file in Word or Google Docs
must see a clean document, not Markdown symbols.
================================================================================
"""
from __future__ import annotations

import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# KairoX document palette — dark slate headings, one accent, grey rules.
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x8A, 0x3B, 0x1E)        # leather brown
SUB = RGBColor(0x55, 0x55, 0x55)
CODE_BG = "F4F4F4"
HEAD_BG = "EDE7E2"
MONO = "Consolas"
BODY = "Calibri"


def _shade(cell_or_para, hex_fill: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_fill)
    target = getattr(cell_or_para, "_tc", None)
    if target is not None:
        target.get_or_add_tcPr().append(el)
    else:
        cell_or_para.paragraph_format.element.get_or_add_pPr().append(el)


_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+\*|`[^`]+`)")


def _write_inline(paragraph, text: str, size: float = 10.5,
                  color: RGBColor = INK, bold: bool = False) -> None:
    """Write text into a paragraph honouring **bold** and `code` markers."""
    for chunk in _INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = MONO
            run.font.size = Pt(size - 0.5)
            run.font.color.rgb = ACCENT
        else:
            run = paragraph.add_run(chunk)
            run.bold = bold
        run.font.size = Pt(size)
        if run.font.color.rgb is None:
            run.font.color.rgb = color


def _write_heading(paragraph, text: str) -> None:
    """Inline markup inside a HEADING, without touching the heading's style.

    `add_heading(text)` writes one literal run, so `## …  (**not** HR)` printed
    the asterisks. This splits the same way `_write_inline` does but sets only
    bold / italic / monospace — the size and colour stay the heading style's.
    """
    for chunk in _INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            paragraph.add_run(chunk[2:-2])
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            paragraph.add_run(chunk[1:-1]).italic = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = MONO
        else:
            paragraph.add_run(chunk)


class DocBuilder:
    """Builds one .docx. Call the methods in document order, then `save()`."""

    def __init__(self) -> None:
        self.doc = Document()
        self._setup()

    # ── chrome ───────────────────────────────────────────────────────────────
    def _setup(self) -> None:
        style = self.doc.styles["Normal"]
        style.font.name = BODY
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15

        for name, size, color in (("Heading 1", 18, ACCENT),
                                  ("Heading 2", 14, INK),
                                  ("Heading 3", 12, INK),
                                  ("Heading 4", 11, SUB)):
            heading = self.doc.styles[name]
            heading.font.name = BODY
            heading.font.size = Pt(size)
            heading.font.bold = True
            heading.font.color.rgb = color

        section = self.doc.sections[0]
        section.left_margin = section.right_margin = Inches(0.85)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    def footer(self, text: str) -> None:
        para = self.doc.sections[0].footer.paragraphs[0]
        para.text = text
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = SUB

    def cover(self, product: str, service: str, kind: str, subtitle: str,
              meta: list[tuple[str, str]]) -> None:
        for _ in range(3):
            self.doc.add_paragraph()
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(product)
        run.font.size = Pt(13)
        run.font.color.rgb = SUB
        run.bold = True

        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(service)
        run.font.size = Pt(34)
        run.font.color.rgb = ACCENT
        run.bold = True

        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(kind)
        run.font.size = Pt(20)
        run.font.color.rgb = INK

        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(14)
        run = para.add_run(subtitle)
        run.font.size = Pt(11)
        run.font.color.rgb = SUB

        self.doc.add_paragraph()
        table = self.doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for key, value in meta:
            row = table.add_row()
            row.cells[0].width = Inches(1.9)
            row.cells[1].width = Inches(4.2)
            _write_inline(row.cells[0].paragraphs[0], f"**{key}**", 9.5)
            _write_inline(row.cells[1].paragraphs[0], value, 9.5)
            _shade(row.cells[0], HEAD_BG)
        self.page_break()

    # ── blocks ───────────────────────────────────────────────────────────────
    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _heading(self, text: str, level: int, space_before: float):
        heading = self.doc.add_heading("", level=level)
        _write_heading(heading, text)
        if space_before:
            heading.paragraph_format.space_before = Pt(space_before)
        return heading

    def h1(self, text: str, page_break: bool = True) -> None:
        if page_break:
            self.page_break()
        self._heading(text, 1, 0)

    def h2(self, text: str) -> None:
        self._heading(text, 2, 14)

    def h3(self, text: str) -> None:
        self._heading(text, 3, 10)

    def h4(self, text: str) -> None:
        self._heading(text, 4, 8)

    def para(self, text: str, size: float = 10.5, bold: bool = False,
             color: RGBColor = INK, space_after: float = 6) -> None:
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(space_after)
        _write_inline(para, text, size, color, bold)

    def bullet(self, text: str, level: int = 0) -> None:
        style = "List Bullet" if level == 0 else "List Bullet 2"
        para = self.doc.add_paragraph(style=style)
        para.paragraph_format.space_after = Pt(2)
        _write_inline(para, text)

    def numbered(self, text: str, level: int = 0) -> None:
        style = "List Number" if level == 0 else "List Number 2"
        para = self.doc.add_paragraph(style=style)
        para.paragraph_format.space_after = Pt(2)
        _write_inline(para, text)

    def code(self, text: str, size: float = 8.5) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.cell(0, 0)
        _shade(cell, CODE_BG)
        cell.paragraphs[0].text = ""
        first = True
        for line in text.rstrip("\n").split("\n"):
            para = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(line)
            run.font.name = MONO
            run.font.size = Pt(size)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def callout(self, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.cell(0, 0)
        _shade(cell, "FFF6E5")
        cell.paragraphs[0].text = ""
        _write_inline(cell.paragraphs[0], text, 10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def table(self, headers: list[str], rows: list[list[str]],
              widths: list[float] | None = None, size: float = 9) -> None:
        if not rows:
            return
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = True
        for index, head in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.paragraphs[0].text = ""
            # An empty header cell is a real thing in a Markdown table (the
            # top-left corner of a label/value grid). Wrapping "" in ** would
            # print four literal asterisks.
            _write_inline(cell.paragraphs[0],
                          f"**{head}**" if str(head).strip() else "", size)
            _shade(cell, HEAD_BG)
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row[:len(headers)]):
                cells[index].paragraphs[0].text = ""
                cells[index].paragraphs[0].paragraph_format.space_after = Pt(1)
                _write_inline(cells[index].paragraphs[0], str(value), size)
        if widths:
            for row in table.rows:
                for index, width in enumerate(widths[:len(headers)]):
                    row.cells[index].width = Inches(width)
        # repeat the header row on every page a long table spills onto
        header = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header.append(repeat)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def rule(self) -> None:
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "CCCCCC")
        border.append(bottom)
        para._p.get_or_add_pPr().append(border)

    def save(self, path) -> None:
        self.doc.save(str(path))


# ─────────────────────────────────────────────────────────────────────────────
# Markdown → Word
# ─────────────────────────────────────────────────────────────────────────────
_TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def render_markdown(builder: DocBuilder, text: str,
                    first_h1_breaks: bool = False) -> None:
    """Render the supported Markdown subset into `builder`."""
    lines = text.replace("\r\n", "\n").split("\n")
    index = 0
    seen_h1 = False
    paragraph: list[str] = []

    def flush() -> None:
        nonlocal paragraph
        if paragraph:
            builder.para(" ".join(paragraph).strip())
            paragraph = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush()
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            index += 1
            builder.code("\n".join(block))
            continue

        if (stripped.startswith("|") and index + 1 < len(lines)
                and _TABLE_SEP.match(lines[index + 1])):
            flush()
            headers = _split_row(stripped)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_row(lines[index]))
                index += 1
            builder.table(headers, rows)
            continue

        if not stripped:
            flush()
            index += 1
            continue

        if stripped.startswith("#### "):
            flush()
            builder.h4(stripped[5:].strip())
        elif stripped.startswith("### "):
            flush()
            builder.h3(stripped[4:].strip())
        elif stripped.startswith("## "):
            flush()
            builder.h2(stripped[3:].strip())
        elif stripped.startswith("# "):
            flush()
            builder.h1(stripped[2:].strip(),
                       page_break=first_h1_breaks and seen_h1)
            seen_h1 = True
        elif stripped.startswith("> "):
            flush()
            builder.callout(stripped[2:].strip())
        elif stripped in ("---", "***", "___"):
            flush()
            builder.rule()
        elif re.match(r"^[-*] ", stripped):
            flush()
            level = 1 if line.startswith(("  -", "  *", "    -")) else 0
            builder.bullet(stripped[2:].strip(), level)
        elif re.match(r"^\d+\. ", stripped):
            flush()
            builder.numbered(re.sub(r"^\d+\.\s*", "", stripped))
        else:
            paragraph.append(stripped)
        index += 1

    flush()
