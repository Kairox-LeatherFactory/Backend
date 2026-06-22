"""
================================================================================
tests/test_procurement_stage2_trigger.py — the Stage-1 → Stage-2 generate trigger
================================================================================

Covers the bridge: ProcurementService.generate_bom_from_submission (exposed at
POST /procurement/submissions/{id}/generate-bom). It consumes a COMPLETE Stage-1
submission into a DRAFT BOM and flips the submission to CONSUMED.

The new contract: the BOM is built from the ORDER + SPEC sheets ALONE — no order/style
is supplied or required. The BOM is anchored on the submission (client_order_id/style_id
NULL); order_qty + the per-size breakdown come from PARSING the order sheet. The
Client→Order→Style→SKU breakdown is materialised only at MD approval and back-linked.

Exercised at the service level against the REAL Beau Geste spec grid (deterministic, no
LLM) plus a synthetic XLSX order sheet so the whole engine runs end-to-end. Storage is
isolated to a tmp dir so the promoted bytes round-trip through core.storage as in prod.
Skipped when the real spec sheet isn't checked out.
================================================================================
"""
import hashlib
import io
import uuid
from pathlib import Path

import pytest
import yaml
from fastapi import HTTPException
from sqlalchemy import select

from app.core.enums import DocumentKind, SpecType, UserRole
from app.core.models import Document
from app.core.storage import get_storage, reset_storage_cache, submission_key
from app.modules.bom.enums import BomStatus
from app.modules.bom.extraction import extract_order
from app.modules.bom.models import Bom, GarmentType, PomDictionary
from app.modules.bom.seed_stage2 import _GARMENT_YAML, _POM_DICT_YAML
from app.modules.bom.service import BomService
from app.modules.clients.models import SKU, Client, ClientOrder, Style
from app.modules.procurement.enums import SubmissionStatus, ValidationStatus
from app.modules.procurement.models import Submission
from app.modules.procurement.service import ProcurementService
from app.modules.procurement.sniffing import XLSX_MIME
from app.modules.users.models import User

DATA = Path(__file__).resolve().parent.parent / "data"
SPEC = DATA / "spec_sheet_1.xlsx"

pytestmark = pytest.mark.skipif(not SPEC.exists(),
                                reason="real Beau Geste spec sheet not present")

# The per-size breakdown the synthetic order sheet declares (Σ = 60 — the workflow doc's
# CARNABY example S:4 M:17 L:23 XL:10 XXL:6, matching the spec grid's five sizes).
ORDER_PER_SIZE = {"S": 4, "M": 17, "L": 23, "XL": 10, "XXL": 6}
ORDER_QTY = sum(ORDER_PER_SIZE.values())


def _order_xlsx_bytes(style="SIDE SUEDE TRACK PANT", color="BLK") -> bytes:
    """A minimal one-block order sheet: S.NO | STYLE | COLOUR | <sizes> | TOTAL."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    sizes = list(ORDER_PER_SIZE)
    ws.append(["S.NO", "STYLE", "COLOUR", *sizes, "TOTAL"])
    ws.append([1, style, color, *[ORDER_PER_SIZE[s] for s in sizes], ORDER_QTY])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def _isolate_storage(tmp_path, monkeypatch):
    """Root local storage at a tmp dir so bytes round-trip without touching the repo
    working tree; reset the cached backend so test + service resolve the same backend."""
    from app.core.config import settings
    monkeypatch.setattr(settings, "local_storage_dir", str(tmp_path))
    reset_storage_cache()
    yield
    reset_storage_cache()


async def _seed_stage2(db) -> None:
    for r in yaml.safe_load(open(_GARMENT_YAML, encoding="utf-8")):
        db.add(GarmentType(code=r["code"], label=r.get("label"),
                           required_poms=r.get("required_poms") or [],
                           area_formula=r.get("area_formula") or {},
                           default_wastage_pct=r.get("default_wastage_pct")))
    await db.flush()
    for r in yaml.safe_load(open(_POM_DICT_YAML, encoding="utf-8")):
        db.add(PomDictionary(language=r["language"], source_term=r["source_term"],
                             pom_code=r["pom_code"], weight=r.get("weight", 1)))
    await db.commit()


async def _dm_user(db) -> User:
    u = User(id=uuid.uuid4(), name="DM", phone="9000000001",
             role=UserRole.DIRECT_MANAGER, password_hash="x", is_active=True)
    db.add(u)
    await db.commit()
    return u


async def _client(db) -> Client:
    client = Client(name="Beau Geste", country="Japan", code="BG", currency="USD")
    db.add(client)
    await db.commit()
    return client


def _promote(sub_id, slot, sha, data, ext=".xlsx"):
    get_storage().put(submission_key(str(sub_id), slot, sha, ext), data)


async def _submission(db, client, *, status=SubmissionStatus.COMPLETE.value) -> Submission:
    """A submission with ACCEPTED order + spec slots and both sets of bytes promoted to
    storage — exactly what Stage 1 leaves behind once both slots are accepted. `client`
    may be None: a submission can be opened before a client is resolved (the client is
    read off the order sheet only at MD approval)."""
    cid = client.id if client else None
    spec_data = SPEC.read_bytes()
    spec_sha = hashlib.sha256(spec_data).hexdigest()
    order_data = _order_xlsx_bytes()
    order_sha = hashlib.sha256(order_data).hexdigest()

    sub = Submission(client_id=cid, status=SubmissionStatus.OPEN.value)
    db.add(sub)
    await db.flush()

    spec_doc = Document(
        client_id=cid, kind=DocumentKind.SPEC_SHEET.value,
        filename="spec_sheet_1.xlsx", mime=XLSX_MIME, sha256=spec_sha,
        submission_id=sub.id, validation_status=ValidationStatus.ACCEPTED.value,
        classified_spec_type=SpecType.MEASUREMENT_GRID.value,
        client_match_code="beau_geste", storage_url=f"file:///stub/{spec_sha}.xlsx",
    )
    order_doc = Document(
        client_id=cid, kind=DocumentKind.ORDER_SHEET.value,
        filename="order_sheet_1.xlsx", mime=XLSX_MIME, sha256=order_sha,
        submission_id=sub.id, validation_status=ValidationStatus.ACCEPTED.value,
        client_match_code="beau_geste", storage_url=f"file:///stub/{order_sha}.xlsx",
    )
    db.add_all([spec_doc, order_doc])
    await db.flush()
    sub.spec_document_id = spec_doc.id
    sub.order_document_id = order_doc.id
    sub.status = status
    await db.commit()
    _promote(sub.id, "spec-sheet", spec_sha, spec_data)
    _promote(sub.id, "order-sheet", order_sha, order_data)
    return sub


# ════════════════════════════════════════════════════════════════════════════
# extract_order — deterministic XLSX parse
# ════════════════════════════════════════════════════════════════════════════
def test_extract_order_xlsx_per_size():
    parsed = extract_order(_order_xlsx_bytes(), "order_sheet_1.xlsx", XLSX_MIME, "beau_geste")
    assert parsed["order_qty"] == ORDER_QTY
    assert parsed["per_size_qty"] == ORDER_PER_SIZE
    assert parsed["style_name"] == "SIDE SUEDE TRACK PANT"
    assert len(parsed["lines"]) == 1
    assert parsed["lines"][0]["color_code"] == "BLK"
    assert parsed["lines"][0]["sizes"] == ORDER_PER_SIZE
    assert parsed["warnings"] == []


def test_extract_order_pdf_via_vision_llm(monkeypatch):
    """A scanned PDF order sheet routes through the Gemini VISION path. We inject a fake
    multimodal model + a fake PDF rasteriser (no API key), mirroring how the procurement
    tests fake the classifier — proving the vision-first wiring end-to-end."""
    import app.modules.procurement.classifier as clf
    import app.modules.procurement.sniffing as sniff
    from app.core.config import settings as cfg

    class _FakeResp:
        content = ('{"order_number":"1579","style_name":"CARNABY","customer_ref":"CR1",'
                   '"currency":"USD","lines":[{"color":"BLACK","sizes":'
                   '{"S":4,"M":17,"L":23,"XL":10,"XXL":6}}]}')

    class _FakeModel:
        def invoke(self, content):           # content is [HumanMessage(text + images)]
            assert isinstance(content, list)
            return _FakeResp()

    monkeypatch.setattr(cfg, "vision_classifier_enabled", True)
    monkeypatch.setattr(cfg, "vision_model", "gemini:fake")
    monkeypatch.setattr(clf, "_init_model", lambda spec: _FakeModel())
    monkeypatch.setattr(sniff, "render_pdf_pages", lambda data, **kw: [b"\x89PNG-fake"])

    parsed = extract_order(b"%PDF-1.4 scanned", "order.pdf", "application/pdf", "beau_geste")
    assert parsed["order_qty"] == ORDER_QTY
    assert parsed["per_size_qty"] == ORDER_PER_SIZE
    assert parsed["style_name"] == "CARNABY"
    assert parsed["order_number"] == "1579"
    assert "order_extracted_by_llm" in parsed["warnings"]


def test_extract_order_pdf_no_model_falls_back(monkeypatch):
    """With vision disabled (no key), a PDF degrades to the pypdf best-effort + a warning
    — never crashes, so Stage 2 still produces a hand-correctable BOM."""
    from app.core.config import settings as cfg
    monkeypatch.setattr(cfg, "vision_classifier_enabled", False)

    parsed = extract_order(b"%PDF-1.4 not-a-real-pdf", "order.pdf", "application/pdf")
    assert parsed["order_qty"] == 0
    assert any("order_sheet_no_text_layer" in w or "order_pdf_unreadable" in w
               for w in parsed["warnings"])


# ════════════════════════════════════════════════════════════════════════════
# generate (Stage-1 → Stage-2)
# ════════════════════════════════════════════════════════════════════════════
async def test_generate_from_submission_happy_path(db):
    await _seed_stage2(db)
    user = await _dm_user(db)
    client = await _client(db)
    sub = await _submission(db, client)

    out = await ProcurementService(db).generate_bom_from_submission(user, sub.id)

    assert out["status"] == "consumed"
    assert out["bom"]["status"] == BomStatus.DRAFT.value
    # the engine ran end-to-end: 14 POMs × 5 sizes extracted deterministically
    assert out["extraction"]["poms"] == 14 * 5
    # order qty + per-size came from PARSING the order sheet (no SKUs existed)
    assert out["order"]["order_qty"] == ORDER_QTY
    assert out["order"]["per_size_qty"] == ORDER_PER_SIZE
    # the configured cost lines were seeded so the BOM is never born empty
    cats = {i["category"] for i in out["bom"]["items"]}
    assert {"manufacturing", "packaging", "fob_charge"} <= cats

    # a DRAFT bom now exists ANCHORED ON THE SUBMISSION — no order/style yet
    bom = (await db.execute(select(Bom).where(Bom.submission_id == sub.id))).scalar_one()
    assert bom.status == BomStatus.DRAFT.value
    assert bom.client_order_id is None and bom.style_id is None
    assert bom.client_id == client.id
    assert bom.order_qty == ORDER_QTY

    # the submission flipped to consumed; the order link is NOT set yet (waits for approval)
    await db.refresh(sub)
    assert sub.status == SubmissionStatus.CONSUMED.value
    assert sub.client_order_id is None


async def test_breakdown_is_materialised_on_approval(db):
    await _seed_stage2(db)
    user = await _dm_user(db)
    client = await _client(db)
    sub = await _submission(db, client)

    out = await ProcurementService(db).generate_bom_from_submission(user, sub.id)
    bom_id = uuid.UUID(out["bom"]["id"])

    bsvc = BomService(db)
    await bsvc.confirm_cutting(user, bom_id)
    await bsvc.approve_bom(user, bom_id)

    # the Client→Order→Style→SKU tree now exists, derived from the order sheet
    order = (await db.execute(
        select(ClientOrder).where(ClientOrder.client_id == client.id))).scalar_one()
    style = (await db.execute(
        select(Style).where(Style.client_order_id == order.id))).scalar_one()
    skus = (await db.execute(select(SKU).where(SKU.style_id == style.id))).scalars().all()
    assert {s.size: s.qty_ordered for s in skus} == ORDER_PER_SIZE
    assert {s.color_code for s in skus} == {"BLK"}
    assert sum(s.qty_ordered for s in skus) == ORDER_QTY

    # the BOM is back-linked to the freshly materialised order/style …
    bom = (await db.execute(select(Bom).where(Bom.id == bom_id))).scalar_one()
    assert bom.client_order_id == order.id
    assert bom.style_id == style.id
    # … and so is the submission
    await db.refresh(sub)
    assert sub.client_order_id == order.id


async def test_double_generate_is_rejected(db):
    await _seed_stage2(db)
    user = await _dm_user(db)
    client = await _client(db)
    sub = await _submission(db, client)
    svc = ProcurementService(db)

    await svc.generate_bom_from_submission(user, sub.id)

    with pytest.raises(HTTPException) as ei:
        await svc.generate_bom_from_submission(user, sub.id)
    assert ei.value.status_code == 409
    assert ei.value.detail["error"] == "submission_already_consumed"


async def test_incomplete_submission_is_rejected(db):
    await _seed_stage2(db)
    user = await _dm_user(db)
    client = await _client(db)
    sub = await _submission(db, client, status=SubmissionStatus.OPEN.value)

    with pytest.raises(HTTPException) as ei:
        await ProcurementService(db).generate_bom_from_submission(user, sub.id)
    assert ei.value.status_code == 409
    assert ei.value.detail["error"] == "submission_not_ready"


# ════════════════════════════════════════════════════════════════════════════
# Regression — review fixes (H1/H2/H3/M1)
# ════════════════════════════════════════════════════════════════════════════
async def test_generate_from_client_less_submission(db):
    """H1: a submission opened with NO client must still generate a DRAFT BOM — the spec
    sheet / pattern / template client_id are nullable now (client resolved at approval).
    Previously this hit a NOT NULL violation on the spec_sheet INSERT."""
    await _seed_stage2(db)
    user = await _dm_user(db)
    sub = await _submission(db, None)            # client_id is None

    out = await ProcurementService(db).generate_bom_from_submission(user, sub.id)

    assert out["status"] == "consumed"
    assert out["bom"]["status"] == BomStatus.DRAFT.value
    bom = (await db.execute(select(Bom).where(Bom.submission_id == sub.id))).scalar_one()
    assert bom.client_id is None


async def test_confirm_cutting_rejected_on_approved_bom(db):
    """H2: confirm_cutting on an already-approved BOM must 409, not silently un-approve it."""
    await _seed_stage2(db)
    user = await _dm_user(db)
    client = await _client(db)
    sub = await _submission(db, client)
    out = await ProcurementService(db).generate_bom_from_submission(user, sub.id)
    bom_id = uuid.UUID(out["bom"]["id"])

    bsvc = BomService(db)
    await bsvc.confirm_cutting(user, bom_id)
    await bsvc.approve_bom(user, bom_id)

    with pytest.raises(HTTPException) as ei:
        await bsvc.confirm_cutting(user, bom_id)
    assert ei.value.status_code == 409
    assert ei.value.detail["error"] == "invalid_state_for_confirmation"


async def test_confirm_cutting_rejected_on_rejected_bom(db):
    """H2: a REJECTED BOM must go through reopen_bom (which bumps revision) — confirm_cutting
    must not jump it straight back to ready_for_review."""
    await _seed_stage2(db)
    user = await _dm_user(db)
    client = await _client(db)
    sub = await _submission(db, client)
    out = await ProcurementService(db).generate_bom_from_submission(user, sub.id)
    bom_id = uuid.UUID(out["bom"]["id"])

    bsvc = BomService(db)
    await bsvc.confirm_cutting(user, bom_id)
    await bsvc.reject_bom(user, bom_id, reason="DCM off")

    with pytest.raises(HTTPException) as ei:
        await bsvc.confirm_cutting(user, bom_id)
    assert ei.value.status_code == 409


async def test_edit_nonmaterial_qty_does_not_reopen_gate(db):
    """H3: editing the quantity of a NON-material line (packaging/manufacturing/FOB) on a
    confirmed BOM must not stamp a DCM source nor bounce the BOM out of review."""
    await _seed_stage2(db)
    user = await _dm_user(db)
    client = await _client(db)
    sub = await _submission(db, client)
    out = await ProcurementService(db).generate_bom_from_submission(user, sub.id)
    bom_id = uuid.UUID(out["bom"]["id"])

    bsvc = BomService(db)
    confirmed = await bsvc.confirm_cutting(user, bom_id)
    assert confirmed["status"] == BomStatus.READY_FOR_REVIEW.value

    # pick a non-DCM cost line (manufacturing/packaging/fob_charge are seeded)
    non_material = next(i for i in out["bom"]["items"]
                        if i["category"] in ("manufacturing", "packaging", "fob_charge"))
    bom = (await db.execute(select(Bom).where(Bom.id == bom_id))).scalar_one()
    res = await bsvc.edit_bom_items(
        user, bom_id, base_revision=bom.revision,
        edits=[{"bom_item_id": non_material["id"], "field": "qty_per_garment", "value": 2}],
    )

    assert res["reconfirm_required"] is False
    bom = (await db.execute(select(Bom).where(Bom.id == bom_id))).scalar_one()
    assert bom.status == BomStatus.READY_FOR_REVIEW.value
    edited = next(i for i in res["recomputed"]["items"] if i["id"] == non_material["id"])
    assert edited["dcm_source"] is None          # no bogus DCM stamp on a cost line


async def test_generate_is_idempotent_on_retry(db):
    """M1: if a prior run created the BOM but didn't persist the `consumed` flip, a retry
    must replay the existing BOM (finishing the transition) — never duplicate it (which
    would hit uq_bom_submission)."""
    await _seed_stage2(db)
    user = await _dm_user(db)
    client = await _client(db)
    sub = await _submission(db, client)
    svc = ProcurementService(db)

    out1 = await svc.generate_bom_from_submission(user, sub.id)
    bom_id = out1["bom"]["id"]

    # simulate the crash window: BOM committed, submission consume NOT persisted
    sub.status = SubmissionStatus.COMPLETE.value
    await db.commit()

    out2 = await svc.generate_bom_from_submission(user, sub.id)
    assert out2.get("replayed") is True
    assert out2["bom"]["id"] == bom_id
    boms = (await db.execute(select(Bom).where(Bom.submission_id == sub.id))).scalars().all()
    assert len(boms) == 1
    await db.refresh(sub)
    assert sub.status == SubmissionStatus.CONSUMED.value
